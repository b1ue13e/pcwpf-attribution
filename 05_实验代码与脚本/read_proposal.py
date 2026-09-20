#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')

try:
    from docx import Document
    print('=== 开始读取Word文档 ===')
    doc = Document('安徽财经大学本科生科研创新基金项目申请书.docx')
    print('=== Word文档内容 ===')
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            print(f'[段落{i}] {text[:80]}')
        if i > 60:
            break
    print('\n=== 表格数量 ===')
    print(f'总表格数: {len(doc.tables)}')
    for i, table in enumerate(doc.tables):
        print(f'\n表{i}: {len(table.rows)}行 x {len(table.columns)}列')
        for j, row in enumerate(table.rows[:2]):
            for k, cell in enumerate(row.cells[:2]):
                cell_text = cell.text[:40]
                print(f'  [{j},{k}]: {cell_text}')
except Exception as e:
    print(f'错误: {e}')
    import traceback
    traceback.print_exc()
