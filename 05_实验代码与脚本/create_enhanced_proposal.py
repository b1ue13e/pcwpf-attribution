#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建完善后的安徽财经大学本科生科研创新基金项目申请书
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--quiet", "--no-deps"])
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def set_chinese_font(run, font_name='宋体', size=12, bold=False):
    """设置中文字体"""
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_paragraph_with_format(doc, text, font_name='宋体', size=12, bold=False, 
                              alignment=WD_PARAGRAPH_ALIGNMENT.LEFT, 
                              first_line_indent=0.74, line_spacing=1.5,
                              space_before=0, space_after=0):
    """添加格式化的段落"""
    para = doc.add_paragraph()
    para.alignment = alignment
    para.paragraph_format.first_line_indent = Cm(first_line_indent) if first_line_indent > 0 else None
    para.paragraph_format.line_spacing = line_spacing
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    
    run = para.add_run(text)
    set_chinese_font(run, font_name, size, bold)
    return para


def add_heading(doc, text, level=1):
    """添加标题"""
    sizes = {1: 16, 2: 14, 3: 12}
    bolds = {1: True, 2: True, 3: True}
    indents = {1: 0, 2: 0.74, 3: 0.74}
    
    return add_paragraph_with_format(
        doc, text, 
        font_name='黑体' if level <= 2 else '宋体',
        size=sizes.get(level, 12),
        bold=bolds.get(level, False),
        first_line_indent=indents.get(level, 0.74),
        line_spacing=1.5
    )


def create_cover_table(doc):
    """创建封面表格"""
    # 标题
    title_para = doc.add_paragraph()
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title_para.add_run('安徽财经大学本科生科研创新基金项目申请书')
    set_chinese_font(run, '黑体', 22, True)
    
    doc.add_paragraph()  # 空行
    
    # 创建封面信息表格
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 设置列宽
    for row in table.rows:
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(8)
    
    cover_data = [
        ('课 题 名 称', '受损SCADA场景下风电功率边缘鲁棒预测研究'),
        ('课 题 类 别', '重点项目'),
        ('课 题 申 请 人', '李俊宇'),
        ('所  在 学  院', '统计与应用数学学院'),
        ('专 业 班 级', '24统计二班'),
        ('课题指导教师', '杜俊涛    副教授'),
    ]
    
    for i, (label, value) in enumerate(cover_data):
        cell0 = table.rows[i].cells[0]
        cell1 = table.rows[i].cells[1]
        
        # 标签单元格
        p0 = cell0.paragraphs[0]
        p0.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run0 = p0.add_run(label)
        set_chinese_font(run0, '宋体', 12, True)
        cell0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        
        # 值单元格
        p1 = cell1.paragraphs[0]
        p1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run1 = p1.add_run(value)
        set_chinese_font(run1, '宋体', 12, False)
        cell1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    
    # 设置表格边框
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(r'<w:tblPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    
    # 添加表格样式
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'''
                <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="single" w:sz="4"/>
                    <w:left w:val="single" w:sz="4"/>
                    <w:bottom w:val="single" w:sz="4"/>
                    <w:right w:val="single" w:sz="4"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
    
    doc.add_paragraph()
    
    # 填表日期
    date_para = doc.add_paragraph()
    date_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = date_para.add_run('填 表 日 期：2026年3月9日')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_paragraph()
    
    # 学校落款
    school_para = doc.add_paragraph()
    school_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = school_para.add_run('安  徽  财  经  大  学')
    set_chinese_font(run, '黑体', 14, True)
    
    doc.add_page_break()


def create_fill_instructions(doc):
    """创建填报说明"""
    add_heading(doc, '填 报 说 明', 1)
    
    instructions = [
        '一、请实事求是、逐条、认真地填写申请书的各项内容。表达应明确、严谨、简洁，字迹工整。',
        '二、外语要同时用原文和中文表达。第一次出现的缩写词须注明全称。',
        '三、申请书为A4本双面印制，页面空格不够时请自行加页，并于左侧装订成册。',
        '四、正式申请时需提交一式一份申请书，需指导教师和所在院系审查、签署意见并加盖印章。',
        '五、同年度每名指导教师指导的项目原则上不超过两项。',
        '六、凡格式不符合要求的申请书，不予接收。',
    ]
    
    for inst in instructions:
        add_paragraph_with_format(doc, inst, first_line_indent=0)
    
    doc.add_page_break()


def create_basic_info(doc):
    """创建基本信息表格"""
    add_heading(doc, '一、基本信息', 1)
    
    # 创建表格
    table = doc.add_table(rows=8, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 合并单元格并填充数据
    info_data = [
        ('项目名称', '受损SCADA场景下风电功率边缘鲁棒预测研究', '项目类别', 'A重点项目'),
        ('申请人', '李俊宇', '院系班级', '统计与应用数学学院 24统计二班'),
        ('学号', '20244353', '联系方式', 'ljylikezmn999@gmail.com'),
        ('项目来源', '学生自主立题', '指导教师', '杜俊涛'),
        ('指导教师联系方式', 'dujuntao@aufe.edu.cn', '', ''),
    ]
    
    # 填充基本信息
    row_idx = 0
    for item in info_data:
        if len(item) == 4:
            table.rows[row_idx].cells[0].text = item[0]
            table.rows[row_idx].cells[1].text = item[1]
            table.rows[row_idx].cells[2].text = item[2]
            table.rows[row_idx].cells[3].text = item[3]
            row_idx += 1
    
    # 课题组成员
    table.rows[5].cells[0].text = '课题组主要成员'
    table.rows[5].cells[1].merge(table.rows[5].cells[3])
    table.rows[5].cells[1].text = '李俊宇（20244353，24统计二班，负责人）'
    
    # 内容摘要
    table.rows[6].cells[0].text = '内容摘要'
    table.rows[6].cells[0].merge(table.rows[6].cells[3])
    table.rows[6].cells[0].text = '（限200字）'
    
    table.rows[7].cells[0].merge(table.rows[7].cells[3])
    abstract = '''本项目聚焦风电场监控与数据采集系统（SCADA）受损场景下的短期功率预测问题，面向风电场控制器侧部署需求，研究在缺失、噪声、重复记录与漂移干扰下仍具备稳定精度、较低计算开销与物理可行性的预测方法。项目拟构建稀疏混合专家模型，引入多尺度分解、动态路由与可行性约束，兼顾预测精度、鲁棒性和边缘部署效率，并通过调度价值指标验证模型应用效果。'''
    table.rows[7].cells[0].text = abstract
    
    # 设置边框
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            from docx.oxml import parse_xml
            tcBorders = parse_xml(r'''
                <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="single" w:sz="4"/>
                    <w:left w:val="single" w:sz="4"/>
                    <w:bottom w:val="single" w:sz="4"/>
                    <w:right w:val="single" w:sz="4"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
    
    doc.add_paragraph()


def create_project_demonstration(doc):
    """创建项目论证部分"""
    add_heading(doc, '二、项目论证', 1)
    
    # 选题依据
    add_heading(doc, '（一）选题依据', 2)
    
    # 1. 研究背景
    add_heading(doc, '1. 研究背景与问题提出', 3)
    
    background_texts = [
        '风电并网规模不断扩大，短期功率预测已不只是统计意义上的误差比较，而会直接影响备用容量安排、弃风控制、平衡成本和调度安全。现有不少研究在清洗后的理想数据集上验证模型，但真实风电场SCADA数据常伴随缺失、重复、异常点、量测漂移和告警混杂等问题，导致模型在实际部署时出现性能塌缩。与此同时，控制器侧部署还受限于批量为1、时延敏感、内存受限等边缘计算条件，难以直接使用参数量大、推理路径固定的深层模型。',
        
        '近年来，风电预测研究主要沿三条路线推进：一是分解重构与混合建模，通过分离趋势、波动和突变提升非平稳序列建模能力；二是基于注意力、图结构与时间序列基础模块的深度模型，提升多变量时序表征能力；三是加入物理约束或不确定性建模，抑制超额定功率等不合理输出。但这三类工作往往分别关注"精度""表示能力"或"物理一致性"，较少将受损SCADA、边缘部署与可行性控制放在同一框架下统筹研究。',
        
        '因此，本项目拟以论文现有工作为基础，围绕受损SCADA场景下的风电功率边缘预测展开研究，既回应数据质量不稳定带来的工程难题，也服务学校在统计建模、机器学习与能源数据分析交叉方向的人才培养需求。'
    ]
    
    for text in background_texts:
        add_paragraph_with_format(doc, text)
    
    # 2. 国内外研究现状
    add_heading(doc, '2. 国内外研究现状', 3)
    
    # 2.1 风电功率预测方法
    add_heading(doc, '（1）风电功率预测方法研究现状', 3)
    
    add_heading(doc, '① 分解重构与混合建模路线', 3)
    route1_texts = [
        '通过分离趋势、波动和突变提升非平稳序列建模能力。早期研究主要基于统计方法如ARIMA模型，随后演变为结合信号分解技术（如变分模态分解VMD、经验模态分解EMD）与机器学习模型的混合方法。',
        'Duan等（2025）提出了针对非平稳信号的分解与元启发式优化框架；Li等（2023）将特征分离融合技术与改进CNN相结合用于综合能源系统多能负荷预测；Xue等（2020）采用VMD分解、KPCA降维与麻雀搜索算法优化LSSVM进行短期风电预测。'
    ]
    for text in route1_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '② 基于深度学习的表征学习方法', 3)
    route2_texts = [
        '近年来，基于注意力机制、图结构和长上下文学习的深度模型在时间序列预测领域取得显著进展。Informer（Zhou等, 2021）通过ProbSparse自注意力机制实现长序列高效预测；Autoformer（Wu等, 2021）引入自相关机制捕捉序列周期性；iTransformer（Liu等, 2024）采用"倒置"架构，将注意力机制应用于变量维度而非时间维度；PatchTST（Nie等, 2023）将时间序列分块处理，实现了与NLP领域的跨域迁移。然而，这些模型大多在清洗后的基准数据集上验证，较少在受损SCADA场景下评估。'
    ]
    for text in route2_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '③ 物理引导与可行性约束方法', 3)
    route3_texts = [
        '功率曲线建模和物理约束引导是提升预测可行性的重要途径。Astolfi等（2015）利用数据挖掘技术分析风电场性能曲线；Martin-Calzada等（2025）采用参数迁移学习优化功率曲线建模中的传感器数据预处理；Liu等（2024）将物理信息嵌入强化学习框架用于极端事件下的概率风电预测。这些方法在训练阶段引入物理约束，但仍依赖于较丰富的辅助状态信息。'
    ]
    for text in route3_texts:
        add_paragraph_with_format(doc, text)
    
    # 2.2 SCADA数据质量
    add_heading(doc, '（2）SCADA数据质量与异常检测研究', 3)
    scada_texts = [
        'SCADA数据质量问题已引起学术界和工业界的广泛关注。Leahy等（2019）系统分析了风电状态监测与可靠性分析中的数据质量问题；Letzgus（2020）研究了变点检测在正常行为模型鲁棒条件监测中的应用；Renstrom等（2020）采用深度自编码器实现风机系统级异常检测；Maldonado-Correa等（2020）通过系统文献综述总结了SCADA数据在风机状态监测中的应用；Vervlimmeren等（2025）提出了基于自编码器和机群中值滤波的SCADA驱动故障预测方法。',
        '这些研究多聚焦于数据清洗和异常检测本身，而非直接面向受损数据条件下的预测任务。'
    ]
    for text in scada_texts:
        add_paragraph_with_format(doc, text)
    
    # 2.3 边缘计算
    add_heading(doc, '（3）边缘计算与控制器侧部署研究', 3)
    edge_texts = [
        '边缘智能在能源系统中的应用日益受到重视。Zhou等（2019）提出了边缘智能的概念框架，强调在边缘设备上实现人工智能推理的重要性；Gooi等（2023）综述了边缘智能在智能电网中的应用；Mwangi等（2024）探讨了软件定义工业物联网边缘网络在下一代海上风电场中的应用。这些研究明确了控制器侧推理的约束条件：批处理大小为1、延迟受限、内存受限、通信受限。'
    ]
    for text in edge_texts:
        add_paragraph_with_format(doc, text)
    
    # 2.4 经济价值
    add_heading(doc, '（4）预测经济价值与调度耦合研究', 3)
    value_texts = [
        '预测误差的经济影响是评估预测方法实用价值的关键。Wang等（2011）研究了风电预测不确定性与机组组合问题；Khaloie等（2020）提出了风电-热电-储能协调报价策略；Yan等（2025）提出了风电预测经济价值的数据驱动评估方法；Houben等（2023）研究了不确定性下的多能源系统微网最优调度。这些研究表明，预测质量的评估不能仅依赖统计指标，还需考虑其对下游决策的实际影响。'
    ]
    for text in value_texts:
        add_paragraph_with_format(doc, text)
    
    # 3. 现有研究的不足
    add_heading(doc, '3. 现有研究的不足与本项目的切入点', 3)
    
    gap_texts = [
        '综上所述，现有研究存在以下三方面不足：',
        '（1）数据侧：多数预测模型在清洗后的数据集上验证，与真实SCADA场景存在差距。风电场SCADA数据常伴随缺失、重复、异常点、量测漂移和告警混杂等问题，现有研究较少系统评估模型在受损数据条件下的性能表现。',
        '（2）部署侧：控制器侧推理受限于批处理大小为1、延迟敏感、内存受限等边缘计算条件，但现有深度预测模型往往参数量大、推理路径固定，难以满足边缘部署需求。',
        '（3）决策侧：预测质量的评估通常独立于下游决策，缺乏统一的预测-调度接口来比较不同预测模型对实际运行的影响。',
        '针对上述不足，本项目拟构建面向受损SCADA场景的边缘鲁棒预测框架，将数据质量控制、边缘部署约束与物理可行性引导统一纳入预测模型设计与评估流程。'
    ]
    for text in gap_texts:
        add_paragraph_with_format(doc, text)
    
    # 添加参考文献
    add_heading(doc, '4. 主要参考文献', 3)
    
    references = [
        '[1] Leahy K, Gallagher C, O\'Donovan P, et al. Issues with Data Quality for Wind Turbine Condition Monitoring and Reliability Analyses[J]. Energies, 2019, 12(2): 201.',
        '[2] Letzgus S. Change-point detection in wind turbine SCADA data for robust condition monitoring with normal behaviour models[J]. Wind Energy Science, 2020, 5: 1375-1397.',
        '[3] Renstrom N, Bangalore P, Highcock E. System-wide anomaly detection in wind turbines using deep autoencoders[J]. Renewable Energy, 2020, 157: 647-659.',
        '[4] Maldonado-Correa J, Martin-Martinez S, Artigao E, et al. Using SCADA Data for Wind Turbine Condition Monitoring: A Systematic Literature Review[J]. Energies, 2020, 13(12): 3132.',
        '[5] Vervlimmeren I, Chesterman X, Verstraeten T, et al. Scalable SCADA-driven failure prediction for offshore wind turbines[J]. Wind Energy Science, 2025, 10: 2615-2637.',
        '[6] Zhou H, Zhang S, Peng J, et al. Informer: Beyond Efficient Transformer for Long Sequence Time-Series Forecasting[C]//AAAI, 2021: 11106-11115.',
        '[7] Wu H, Xu J, Wang J, et al. Autoformer: Decomposition Transformers with Auto-Correlation for Long-Term Series Forecasting[C]//NeurIPS, 2021: 22419-22430.',
        '[8] Liu Y, Hu T, Zhang H, et al. iTransformer: Inverted Transformers Are Effective for Time Series Forecasting[C]//ICLR, 2024.',
        '[9] Nie Y, Nguyen N H, Sinthong P, et al. A Time Series is Worth 64 Words: Long-term Forecasting with Transformers[C]//ICLR, 2023.',
        '[10] Wang S, Wu H, Shi X, et al. TimeMixer: Decomposable Multiscale Mixing for Time Series Forecasting[C]//ICLR, 2024.',
        '[11] Astolfi D, Castellani F, Garinei A, et al. Data mining techniques for performance analysis of onshore wind farms[J]. Applied Energy, 2015, 148: 220-233.',
        '[12] Martin-Calzada P, Martin Sanchez P, Rodriguez-Sanchez F J, et al. Optimized Sensor Data Preprocessing Using Parameter-Transfer Learning for Wind Turbine Power Curve Modeling[J]. Sensors, 2025, 25(17): 5329.',
        '[13] Liu Y, Wang J, Liu L. Physics-informed reinforcement learning for probabilistic wind power forecasting under extreme events[J]. Applied Energy, 2024, 376: 124068.',
        '[14] Zhou Z, Chen X, Li E, et al. Edge Intelligence: Paving the Last Mile of Artificial Intelligence With Edge Computing[J]. Proceedings of the IEEE, 2019, 107(8): 1738-1762.',
        '[15] Gooi H B, Wang Z, Tang Y. Edge Intelligence for Smart Grid: A Survey of Applications, Techniques and Challenges[J]. CSEE Journal of Power and Energy Systems, 2023, 9(5): 1623-1640.',
        '[16] Wang J, Botterud A, Bessa R, et al. Wind power forecasting uncertainty and unit commitment[J]. Applied Energy, 2011, 88(11): 4014-4023.',
        '[17] Khaloie H, Abdollahi A, Shafie-khah M, et al. Coordinated wind-thermal-energy storage offering strategy in energy and spinning reserve markets[J]. Applied Energy, 2020, 259: 114168.',
        '[18] Yan J, Li X, Wang H, et al. The economic value of wind power forecasting: a data-driven method and its application in various scenarios[J]. Energy, 2025, 340: 139151.',
        '[19] Houben N, Cosic A, Stadler M, et al. Optimal dispatch of a multi-energy system microgrid under uncertainty[J]. Applied Energy, 2023, 337: 120913.',
        '[20] Verdone A, Panella M, De Santis E, et al. A review of solar and wind energy forecasting[J]. Applied Energy, 2025, 392: 126016.',
    ]
    
    for ref in references:
        add_paragraph_with_format(doc, ref, first_line_indent=0)


def create_research_content(doc):
    """创建研究内容部分"""
    add_heading(doc, '（二）研究内容', 2)
    
    # 研究目的
    add_heading(doc, '1. 研究目的', 3)
    
    purpose_texts = [
        '本项目的核心研究目的是：在风电场SCADA数据存在缺失、噪声、漂移等质量问题的真实场景下，构建一个满足边缘部署约束（低延迟、低内存占用）且具有物理可行性的短期风电功率预测方法，并通过统一的预测-调度评估接口验证其对实际运行的经济价值。',
        '具体而言，研究目的包括以下四个层面：',
        '（1）数据层面：建立系统化的受损SCADA场景刻画与评估方法，量化不同数据质量问题对预测性能的影响程度。',
        '（2）模型层面：设计轻量化的稀疏混合专家预测架构，通过动态路由机制在保证预测精度的同时满足边缘推理的时延与内存约束。',
        '（3）约束层面：构建训练期可行性引导机制，将风电功率物理边界约束（额定功率、切入/切出风速等）嵌入模型学习过程，抑制不合理输出。',
        '（4）评估层面：建立预测-调度耦合评估框架，除统计精度指标外，从调度成本、弃风量、后备替代频率等角度评价预测方法的实际应用价值。'
    ]
    for text in purpose_texts:
        add_paragraph_with_format(doc, text)
    
    # 研究意义
    add_heading(doc, '2. 研究意义', 3)
    
    add_heading(doc, '（1）理论意义', 3)
    theory_texts = [
        '① 丰富风电预测理论的适用边界。现有预测理论多基于理想数据假设，本项目将预测理论拓展至受损数据场景，为真实风电场环境下的预测建模提供理论参考。',
        '② 拓展边缘智能在时间序列预测领域的应用范式。通过稀疏混合专家架构与条件跳跃机制的设计，为资源受限场景下的深度学习模型部署提供新的技术路径。',
        '③ 深化物理信息嵌入机器学习的方法论。本项目提出的可行性引导训练机制，为将领域物理知识融入数据驱动模型提供了一种可操作的实现方式。'
    ]
    for text in theory_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '（2）实践意义', 3)
    practice_texts = [
        '① 提升风电场运行管理的精细化水平。研究成果可直接应用于风电场控制器侧的功率预测模块，在数据质量不佳的条件下仍能提供可靠的预测支撑。',
        '② 降低预测系统的部署与运维成本。轻量化的模型设计使得预测算法可在现有风电场控制设备上运行，无需额外采购高性能计算设备。',
        '③ 服务区域新能源消纳与电网安全。准确的功率预测有助于降低风电并网带来的不确定性，减少备用容量需求和弃风损失，支撑新型电力系统建设。',
        '④ 促进学科交叉人才培养。本项目融合统计学、机器学习、能源系统等多个学科领域，有助于培养具有跨学科背景的复合型人才。'
    ]
    for text in practice_texts:
        add_paragraph_with_format(doc, text)
    
    # 主要研究内容
    add_heading(doc, '3. 主要研究内容', 3)
    
    intro = '本项目以风电场SCADA时间序列为研究对象，围绕"受损数据输入"与"边缘部署约束"两大核心挑战，开展以下四个方面的研究：'
    add_paragraph_with_format(doc, intro)
    
    # 研究内容1
    add_heading(doc, '（1）受损SCADA场景刻画与数据预处理研究', 3)
    content1_texts = [
        '系统分析风电场SCADA数据常见的质量问题，包括：',
        '· 数据缺失：通信中断、传感器故障导致的间歇性缺失',
        '· 量测噪声：电磁干扰、湍流感应抖动导致的随机误差',
        '· 重复记录：数据采集系统异常导致的重复条目',
        '· 量测漂移：传感器校准偏移导致的系统性偏差',
        '· 异常扰动：设备故障、限功率运行导致的离群点',
        '构建涵盖上述问题的受损SCADA场景库，设计统一的数据预处理流程（缺失值插补、异常检测、标准化处理），建立数据质量评估指标体系。'
    ]
    for text in content1_texts:
        add_paragraph_with_format(doc, text)
    
    # 研究内容2
    add_heading(doc, '（2）基于稀疏混合专家的轻量预测模型研究', 3)
    content2_texts = [
        '设计Ultra-LSNT（Ultra-Lightweight Sparse Network for Time-series）预测架构，核心创新点包括：',
        '① 多尺度分解嵌入层：采用移动平均分解将输入序列分离为趋势项与季节项，通过并行一维卷积提取多尺度特征，实现不同频率模式的自适应表征。',
        '② 稀疏Top-K路由机制：引入混合专家（Mixture of Experts, MoE）架构，通过Top-K路由机制为每个时间步选择最相关的K个专家网络（K远小于专家总数N），将计算复杂度从O(N·L·d²)降低至O(K·L·d²)。',
        '③ 跳跃门控条件推理：设计Jump Gate机制，对平稳时段的表征进行条件跳过，进一步减少推理计算量，使模型能够根据输入复杂度动态调整计算资源分配。',
        '④ 负载均衡正则化：引入辅助损失函数促进专家负载均衡，避免路由崩溃问题，提升模型训练稳定性。'
    ]
    for text in content2_texts:
        add_paragraph_with_format(doc, text)
    
    # 研究内容3
    add_heading(doc, '（3）可行性引导的训练机制研究', 3)
    content3_texts = [
        '针对风电功率预测中常见的超额定功率、负功率等不合理输出问题，构建FG-MoE（Feasibility-Guided MoE）训练框架：',
        '① 物理包络建模：基于风机空气动力学功率转换原理，构建由切入风速、额定风速、切出风速、额定功率等参数确定的可行功率包络。',
        '② 部署一致性代理：采用最后观测风速作为未来风速的代理（因未来风速在推理时不可获取），保证训练阶段使用的可行性约束信息与推理阶段一致。',
        '③ 训练期约束嵌入：在损失函数中引入物理可行性惩罚项，对超出物理包络、违反功率边界、在切出区域非零输出等情况进行惩罚，引导模型学习符合物理规律的预测行为。',
        '④ 与后处理裁剪的对比分析：系统对比训练期约束嵌入（FG-MoE）与推理期后处理裁剪（Clip-only）两种策略在预测精度与可行性之间的权衡关系。'
    ]
    for text in content3_texts:
        add_paragraph_with_format(doc, text)
    
    # 研究内容4
    add_heading(doc, '（4）预测-调度耦合评估方法研究', 3)
    content4_texts = [
        '建立统一的预测-调度评估接口，从实际应用价值角度评价预测方法：',
        '① 统计精度评估：采用R²、RMSE、MAE等指标评估预测精度，并在不同噪声水平（σ_eff ∈ {0.00, 0.15, 0.30, 0.45, 0.60}）下进行鲁棒性切片评估。',
        '② 边缘性能评估：测量模型的批处理-1推理延迟（CPU/ARM平台）、内存占用、参数量等指标，验证其满足边缘部署约束的能力。',
        '③ 安全合规评估：计算超额定功率发生率、切入/切出区域违规率等安全性指标，评估预测输出的物理可行性。',
        '④ 调度价值评估：构建日前调度筛选模型，以120天滚动窗口评估各预测方法在运营成本、弃风量、后备替代频率等方面的表现，验证预测误差向决策成本的传导效应。'
    ]
    for text in content4_texts:
        add_paragraph_with_format(doc, text)
    
    # 技术路线
    add_heading(doc, '4. 研究方案与技术路线', 3)
    
    add_heading(doc, '（1）技术路线图', 3)
    route_texts = [
        '本项目技术路线如下：',
        '数据层：SCADA数据获取 → 质量分析 → 受损场景注入 → 预处理标准化',
        '模型层：多尺度分解 → 稀疏路由 → 条件跳跃 → 预测输出',
        '约束层：物理包络计算 → 可行性惩罚 → 联合优化',
        '评估层：统计指标 → 边缘性能 → 安全合规 → 调度价值'
    ]
    for text in route_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '（2）实验设计', 3)
    exp_texts = [
        '① 数据集：主要数据集为新疆某风电场SCADA数据（15分钟采样，约12万条记录）；迁移验证使用美国风电数据、空气质量数据、GEFCom负荷数据。',
        '② 对比基线：统计方法（ARIMA、Persistence）、机器学习方法（LightGBM、SVR）、深度学习方法（DLinear、TimeMixer、iTransformer）、长上下文基线（Chronos、Moirai）。',
        '③ 主要实验：鲁棒性实验（在不同噪声、漂移、量化水平下评估模型性能）、消融实验（验证稀疏路由、跳跃门控、可行性引导各组件的贡献）、边缘性能实验（在CPU和ARM平台测量推理延迟与内存占用）、调度价值实验（在统一调度接口下比较各模型的实际应用价值）。'
    ]
    for text in exp_texts:
        add_paragraph_with_format(doc, text)


def create_progress_plan(doc):
    """创建研究进度安排"""
    add_heading(doc, '（三）研究进度安排', 2)
    
    schedules = [
        '第一阶段（2026年3月-2026年6月）：文献调研与数据准备。完成相关文献的系统梳理，建立受损SCADA场景库，完成数据预处理流程开发。',
        '第二阶段（2026年7月-2026年12月）：模型设计与实现。完成Ultra-LSNT模型架构设计与代码实现，开展初步实验验证。',
        '第三阶段（2027年1月-2027年6月）：可行性引导机制研究。完成FG-MoE训练机制设计与实现，开展对比实验分析。',
        '第四阶段（2027年7月-2027年12月）：预测-调度耦合评估。完成统一评估接口开发，开展调度价值实验，整理研究成果。',
        '第五阶段（2028年1月-2028年3月）：论文撰写与成果总结。完成学术论文撰写，准备项目结题材料。'
    ]
    
    for schedule in schedules:
        add_paragraph_with_format(doc, schedule)


def create_existing_work(doc):
    """创建已有研究基础"""
    add_heading(doc, '（四）已有研究基础', 2)
    
    add_heading(doc, '1. 模型架构设计与实现', 3)
    work1_texts = [
        '已完成Ultra-LSNT预测模型的设计与实现，该模型融合了多尺度分解、稀疏Top-K路由、跳跃门控等核心组件。模型在PyTorch框架下实现，支持端到端训练与推理。',
        '核心代码结构：分解嵌入层实现趋势-季节分离与多尺度卷积特征提取；稀疏路由模块实现Top-K专家选择与负载均衡；跳跃门控模块实现条件推理路径跳过；可行性约束层实现物理包络计算与训练期惩罚。'
    ]
    for text in work1_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '2. 实验验证与结果分析', 3)
    work2_texts = [
        '已完成以下主要实验：',
        '（1）鲁棒性验证实验：在受损SCADA场景下，Ultra-LSNT在严重噪声水平（σ_eff=0.60）下仍保持R²=0.7863的正向预测能力，而对比方法DLinear在该条件下R²降至-0.9903。',
        '（2）边缘性能评估：在批处理-1推理条件下，Ultra-LSNT-Lite优化版本的CPU延迟为1.171ms，ARM延迟为5.007ms，内存占用仅1.33MiB，满足边缘部署约束。',
        '（3）可行性引导效果验证：FG-MoE在保持清洁条件R²=0.8747、严重条件R²=0.8286的同时，将超额定发生率降至0.0092%，显著优于后处理裁剪策略。',
        '（4）调度价值评估：在120天滚动调度评估中，Ultra-LSNT相比DLinear基线显著降低了运营成本和弃风量。'
    ]
    for text in work2_texts:
        add_paragraph_with_format(doc, text)
    
    add_heading(doc, '3. 论文与代码成果', 3)
    work3_texts = [
        '相关研究成果已整理为学术论文，投稿至Applied Energy期刊。论文系统阐述了受损SCADA场景下的边缘鲁棒预测方法。',
        '代码实现已开源（GitHub: https://github.com/b1ue13e/Ultra-LSNT），包含数据预处理脚本、模型实现、训练与评估代码，支持研究结果的复现。'
    ]
    for text in work3_texts:
        add_paragraph_with_format(doc, text)


def create_innovation(doc):
    """创建创新点与难点"""
    add_heading(doc, '（五）创新点与重难点', 2)
    
    add_heading(doc, '1. 关键科学问题', 3)
    problems = [
        '（1）受损数据条件下的预测稳定性问题：如何在SCADA数据存在缺失、噪声、漂移等质量问题的条件下，保持预测模型的泛化稳定性，避免性能塌缩。',
        '（2）资源约束条件下的模型效率问题：如何在边缘设备有限的计算资源和内存预算下，实现满足实时性要求的预测推理。',
        '（3）物理约束与数据驱动的融合问题：如何将物理可行性约束有效嵌入数据驱动的模型训练过程，在不显著牺牲预测精度的前提下提升输出合理性。'
    ]
    for p in problems:
        add_paragraph_with_format(doc, p)
    
    add_heading(doc, '2. 主要创新点', 3)
    innovations = [
        '（1）研究视角创新：首次将受损SCADA数据、边缘部署约束、物理可行性控制三类实际约束纳入统一的预测研究框架，填补了现有研究主要基于理想数据假设、较少考虑部署约束的空白。',
        '（2）模型架构创新：提出Ultra-LSNT轻量稀疏预测架构，通过Top-K路由与跳跃门控机制实现计算资源的动态分配，在保持预测精度的同时满足边缘推理的延迟与内存约束。',
        '（3）训练机制创新：提出FG-MoE可行性引导训练机制，采用部署一致性的风速代理，将物理包络约束嵌入训练期优化，避免了后处理裁剪带来的精度损失。',
        '（4）评估方法创新：建立预测-调度耦合评估框架，除统计精度外，从运营成本、弃风量、后备替代频率等角度评价预测方法的实际应用价值。'
    ]
    for inn in innovations:
        add_paragraph_with_format(doc, inn)


def create_budget(doc):
    """创建经费预算"""
    add_heading(doc, '三、经费预算', 1)
    
    budget_items = [
        ('资料费', '购买相关书籍、期刊、数据库访问权限等', '2000'),
        ('实验材料费', '计算资源租赁、云服务费用等', '3000'),
        ('设备费', '移动硬盘、存储设备等', '1000'),
        ('差旅费', '参加学术会议、实地调研等', '2000'),
        ('论文版面费', '学术论文发表费用', '2000'),
        ('其他费用', '学术交流、资料打印等', '1000'),
        ('合计', '', '11000'),
    ]
    
    # 创建表格
    table = doc.add_table(rows=len(budget_items) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 表头
    headers = ['序号', '支出科目', '计算依据及理由', '金额（元）']
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(h)
        set_chinese_font(run, '黑体', 11, True)
    
    # 数据行
    for i, (item, reason, amount) in enumerate(budget_items, 1):
        row = table.rows[i]
        if item == '合计':
            row.cells[0].text = ''
            row.cells[1].text = item
            row.cells[1].merge(row.cells[2])
            row.cells[2].text = amount
        else:
            row.cells[0].text = str(i)
            row.cells[1].text = item
            row.cells[2].text = reason
            row.cells[3].text = amount
    
    # 设置边框
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            from docx.oxml import parse_xml
            tcBorders = parse_xml(r'''
                <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="single" w:sz="4"/>
                    <w:left w:val="single" w:sz="4"/>
                    <w:bottom w:val="single" w:sz="4"/>
                    <w:right w:val="single" w:sz="4"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)


def create_signatures(doc):
    """创建签章部分"""
    add_heading(doc, '四、指导教师意见', 1)
    
    opinion_text = '''该项目聚焦风电场SCADA数据受损场景下的边缘鲁棒预测问题，选题具有重要的理论意义和实践价值。申请人已具备较好的数学基础和编程能力，前期已完成了相关模型的设计与初步实验验证。研究方案可行，技术路线清晰，预期成果明确。同意指导该项目。'''
    
    add_paragraph_with_format(doc, opinion_text)
    
    # 签名区域
    doc.add_paragraph()
    doc.add_paragraph()
    
    sign_para = doc.add_paragraph()
    sign_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = sign_para.add_run('指导教师签名：_______________')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_paragraph()
    
    date_para = doc.add_paragraph()
    date_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = date_para.add_run('年   月   日')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_page_break()
    
    # 院系意见
    add_heading(doc, '五、院系审核意见', 1)
    
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    dept_para = doc.add_paragraph()
    dept_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = dept_para.add_run('单位盖章：')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_paragraph()
    
    date_para2 = doc.add_paragraph()
    date_para2.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = date_para2.add_run('年   月   日')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_page_break()
    
    # 学校意见
    add_heading(doc, '六、学校审批意见', 1)
    
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    school_para = doc.add_paragraph()
    school_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = school_para.add_run('单位盖章：')
    set_chinese_font(run, '宋体', 12)
    
    doc.add_paragraph()
    
    date_para3 = doc.add_paragraph()
    date_para3.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = date_para3.add_run('年   月   日')
    set_chinese_font(run, '宋体', 12)


def main():
    """主函数"""
    print("开始创建完善后的申请书...")
    
    # 创建文档
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(12)
    
    # 创建各部分内容
    create_cover_table(doc)
    create_fill_instructions(doc)
    create_basic_info(doc)
    create_project_demonstration(doc)
    create_research_content(doc)
    create_progress_plan(doc)
    create_existing_work(doc)
    create_innovation(doc)
    create_budget(doc)
    create_signatures(doc)
    
    # 保存文档
    output_path = '安徽财经大学本科生科研创新基金项目申请书_完善版.docx'
    doc.save(output_path)
    
    print(f"申请书已完善并保存至: {output_path}")
    print("主要完善内容：")
    print("1. 封面表格已规范对齐")
    print("2. 选题依据已扩充，包含完整的文献综述（分四条路线详述）")
    print("3. 添加了20条主要参考文献")
    print("4. 研究内容已细化为：研究目的、研究意义、主要研究内容、技术路线、实验设计")
    print("5. 增加了已完成的研究工作介绍")
    print("6. 全文已设置首行缩进2字符、1.5倍行距")


if __name__ == '__main__':
    main()
